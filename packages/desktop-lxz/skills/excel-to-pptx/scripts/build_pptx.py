"""
PPTX Builder 鈥?excel-to-pptx skill
鏍规嵁 Layout Plan JSON 鐢熸垚 PPTX 鏂囦欢銆?
鏀寔 per-page theme锛堟潵鑷笉鍚?PPTX 妯℃澘锛夈€佸鍥剧墖甯冨眬鍜屽浘鐗囦笅鏍囥€?

Usage:
  python build_pptx.py <layout_plan.json> --output <output.pptx> [--theme default.json]
"""

import argparse
import json
import sys
import os
import math

from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
from lxml import etree

EMU_PER_CM = 360000
SLIDE_W_CM = 25.4
SLIDE_H_CM = 14.2875

NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'


def cm(val):
    return Emu(int(val * EMU_PER_CM))


def hex_to_rgb(hex_str):
    """Convert '#RRGGBB' to RGBColor. Returns None for scheme colors."""
    if not hex_str or not isinstance(hex_str, str):
        return None
    hex_str = hex_str.lstrip('#')
    if len(hex_str) != 6:
        return None
    try:
        return RGBColor(int(hex_str[0:2], 16), int(hex_str[2:4], 16), int(hex_str[4:6], 16))
    except Exception:
        return None


def make_theme(overrides=None):
    """Create a theme dict with optional overrides."""
    theme = {
        "title_color": "#1B365D",
        "title_font": "Microsoft YaHei",
        "header_bg": "#1B365D",
        "header_fg": "#FFFFFF",
        "header_font": "Microsoft YaHei",
        "row_even": "#F5F5F0",
        "row_odd": "#FFFFFF",
        "body_font": "Microsoft YaHei",
        "border_color": "#000000",
        "accent": "#1B365D",
    }
    if overrides:
        for k, v in overrides.items():
            theme[k] = v
    rgb_theme = {}
    for k, v in theme.items():
        if isinstance(v, str) and v.startswith("#"):
            rgb_theme[k] = hex_to_rgb(v) or v
        else:
            rgb_theme[k] = v
    return rgb_theme


def ensure_rgb(value, fallback):
    if isinstance(value, RGBColor):
        return value
    return hex_to_rgb(value) or fallback


def set_cell_bg(cell, rgb_color):
    """Set cell background color."""
    tcPr = cell._tc.get_or_add_tcPr()
    for old_fill in tcPr.findall(f'{{{NS_A}}}solidFill'):
        tcPr.remove(old_fill)
    if rgb_color is None:
        return
    solidFill = etree.SubElement(tcPr, f'{{{NS_A}}}solidFill')
    srgbClr = etree.SubElement(solidFill, f'{{{NS_A}}}srgbClr')
    if isinstance(rgb_color, RGBColor):
        srgbClr.set('val', str(rgb_color))
    else:
        srgbClr.set('val', str(rgb_color).lstrip('#'))


def set_cell_border(cell, color="000000", width="6350"):
    """Apply a visible border to all cell edges.

    python-pptx does not expose table cell borders as a high-level API, so this
    writes DrawingML line elements directly. Width 6350 EMUs is approximately
    0.5 pt, clear enough for settlement tables without looking heavy.
    """
    tcPr = cell._tc.get_or_add_tcPr()
    for edge in ("lnL", "lnR", "lnT", "lnB"):
        old = tcPr.find(f'{{{NS_A}}}{edge}')
        if old is not None:
            tcPr.remove(old)
        ln = etree.SubElement(tcPr, f'{{{NS_A}}}{edge}')
        ln.set('w', str(width))
        solidFill = etree.SubElement(ln, f'{{{NS_A}}}solidFill')
        srgbClr = etree.SubElement(solidFill, f'{{{NS_A}}}srgbClr')
        srgbClr.set('val', str(color).lstrip('#'))
        prstDash = etree.SubElement(ln, f'{{{NS_A}}}prstDash')
        prstDash.set('val', 'solid')


def set_cell_text(cell, text, font_name, font_size, bold=False, color=None, align=PP_ALIGN.LEFT, max_chars=100):
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    p.alignment = align
    if text is None:
        text = ""
    text = str(text)
    if max_chars and len(text) > max_chars:
        text = text[:max_chars - 1] + "…"
    run = p.add_run()
    run.text = text
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    cell.margin_left = Emu(18000)
    cell.margin_right = Emu(18000)
    cell.margin_top = Emu(9000)
    cell.margin_bottom = Emu(9000)


def add_text(slide, text, x, y, w, h, font_name="Microsoft YaHei", font_size=10,
             color=None, bold=False, align=PP_ALIGN.LEFT):
    tx_box = slide.shapes.add_textbox(cm(x), cm(y), cm(w), cm(h))
    tf = tx_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text or ""
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = color
    return tx_box


def get_page_theme(page, global_theme):
    """Merge page-level theme overrides with global theme."""
    if not page.get("theme"):
        return global_theme
    return make_theme(page["theme"])


def get_page_dims(page, default_w, default_h):
    """Get slide dimensions for a page."""
    dims = page.get("slide_dims")
    if dims:
        return dims["w"], dims["h"]
    return default_w, default_h


def get_blank_layout(prs):
    """Return the best blank-like layout from the active presentation/template."""
    layouts = list(prs.slide_layouts)
    if not layouts:
        raise RuntimeError("PPTX contains no slide layouts")
    for layout in layouts:
        lname = (getattr(layout, "name", "") or "").lower()
        if "blank" in lname or "空白" in lname:
            return layout
    return min(layouts, key=lambda layout: len(layout.placeholders))


def remove_generated_placeholders(slide):
    """Remove editable placeholders cloned from a template layout.

    Some user PPTX templates expose title placeholders such as "点击此处添加标题"
    even when their visual design is otherwise correct. Generated slides add their
    own title/table/image shapes, so cloned placeholders should be removed while
    preserving master/layout background artwork and theme information.
    """
    for shp in list(slide.shapes):
        if getattr(shp, "is_placeholder", False):
            element = shp._element
            parent = element.getparent()
            if parent is not None:
                parent.remove(element)


def add_blank_slide(prs):
    """Add a generated slide using the current presentation's own layout set."""
    slide = prs.slides.add_slide(get_blank_layout(prs))
    remove_generated_placeholders(slide)
    return slide


def clear_template_slides(prs):
    """Remove sample slides from a template while preserving masters, layouts and theme."""
    slide_id_list = prs.slides._sldIdLst
    for slide_id in list(slide_id_list):
        rel_id = slide_id.rId
        prs.part.drop_rel(rel_id)
        slide_id_list.remove(slide_id)


def get_plan_canvas_dims(layout_plan, fallback_w, fallback_h):
    """Return the deterministic layout canvas stored in the layout plan."""
    canvas = layout_plan.get("layout_canvas") or {}
    try:
        if canvas.get("w") and canvas.get("h"):
            return float(canvas["w"]), float(canvas["h"])
    except Exception:
        pass
    for page in layout_plan.get("pages", []):
        dims = page.get("slide_dims") or {}
        try:
            if dims.get("w") and dims.get("h"):
                return float(dims["w"]), float(dims["h"])
        except Exception:
            continue
    return fallback_w, fallback_h


def scale_shape_geometry(shape, sx, sy):
    """Scale template/master/layout shape geometry to the locked layout canvas.

    This preserves the user's template artwork after the presentation canvas is
    normalized to the built-in layout size. Generated tables and images are then
    placed with the same coordinates as the default output.
    """
    for attr, factor in (("left", sx), ("width", sx), ("top", sy), ("height", sy)):
        try:
            value = getattr(shape, attr)
            if value is not None:
                setattr(shape, attr, int(value * factor))
        except Exception:
            pass


def scale_shape_collection(shapes, sx, sy):
    for shape in list(shapes):
        scale_shape_geometry(shape, sx, sy)


def fit_template_to_layout_canvas(prs, target_w_cm, target_h_cm):
    """Resize a user template to the canonical layout canvas without losing its style.

    The generated content layout must stay identical to the default output. When
    a user template has a different physical slide size, scale its master/layout
    artwork into the canonical canvas, then set the presentation dimensions to
    that canvas. Template colors, fonts, masters and layouts remain available.
    """
    old_w = int(prs.slide_width)
    old_h = int(prs.slide_height)
    new_w = int(cm(target_w_cm))
    new_h = int(cm(target_h_cm))
    if old_w <= 0 or old_h <= 0:
        prs.slide_width = Emu(new_w)
        prs.slide_height = Emu(new_h)
        return
    if abs(old_w - new_w) <= 1000 and abs(old_h - new_h) <= 1000:
        return

    sx = new_w / old_w
    sy = new_h / old_h
    for master in prs.slide_masters:
        scale_shape_collection(master.shapes, sx, sy)
    for layout in prs.slide_layouts:
        scale_shape_collection(layout.shapes, sx, sy)
    for slide in prs.slides:
        scale_shape_collection(slide.shapes, sx, sy)
    prs.slide_width = Emu(new_w)
    prs.slide_height = Emu(new_h)


def build_cover(prs, page, theme, sw, sh):
    slide = add_blank_slide(prs)
    tc_rgb = ensure_rgb(theme.get("title_color"), RGBColor(0x1B, 0x36, 0x5D))
    tf_name = theme.get("title_font", "Microsoft YaHei")

    title_h = 2.2
    title_box = add_text(slide, page["title"], 2, (sh - title_h) / 2, sw - 4, title_h,
                         tf_name, 36, tc_rgb, True, PP_ALIGN.CENTER)
    title_box.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE

    accent = ensure_rgb(theme.get("accent"), RGBColor(0x1B, 0x36, 0x5D))
    line_y = (sh - title_h) / 2 + title_h + 0.25
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cm(sw * 0.3), cm(line_y), cm(sw * 0.4), cm(0.08))
    line.fill.solid()
    line.fill.fore_color.rgb = accent
    line.line.fill.background()
    return slide


def build_toc(prs, page, theme, sw, sh):
    slide = add_blank_slide(prs)
    tc_rgb = ensure_rgb(theme.get("title_color"), RGBColor(0x1B, 0x36, 0x5D))
    tf_name = theme.get("title_font", "Microsoft YaHei")
    add_text(slide, "鐩綍", 1.5, 0.8, 10, 1.2, tf_name, 28, tc_rgb, True)

    y = 3.0
    for item in page.get("sheets", []):
        tmpl_info = f" [{item.get('template', '')}]" if item.get('template') else ""
        add_text(slide, f"{item['name']}{tmpl_info}", 2.5, y, sw - 5, 0.8,
                 theme.get("body_font", tf_name), 14, tc_rgb)
        y += 1.0
    return slide


def build_summary(prs, page, theme, sw, sh):
    slide = add_blank_slide(prs)
    tc_rgb = ensure_rgb(theme.get("title_color"), RGBColor(0x1B, 0x36, 0x5D))
    tf_name = theme.get("title_font", "Microsoft YaHei")
    add_text(slide, page.get("name", "汇总"), 1.5, 0.8, sw - 3, 1.2, tf_name, 24, tc_rgb, True)

    headers = page.get("headers", [])
    data = page.get("data", [])
    col_indices = select_summary_columns(headers, data)
    if data and col_indices:
        y = 3.0
        for row in data[:10]:
            parts = []
            for ci in col_indices:
                val = row[ci] if ci < len(row) else None
                if ci < len(headers) and headers[ci] and val is not None:
                    parts.append(f"{headers[ci]}: {str(val)[:30]}")
            text = "  |  ".join(parts)
            add_text(slide, text, 2, y, sw - 4, 0.6,
                     theme.get("body_font", tf_name), 9, RGBColor(0x33, 0x33, 0x33))
            y += 0.6
    return slide


def format_summary_items(summary_context):
    if not summary_context:
        return []
    values = summary_context.get("values") or {}
    if isinstance(values, dict) and values:
        return [(str(k), v) for k, v in values.items() if v is not None and str(v).strip() != ""]
    headers = summary_context.get("headers", [])
    row = summary_context.get("row", [])
    items = []
    for i, value in enumerate(row):
        if value is None or str(value).strip() == "":
            continue
        key = headers[i] if i < len(headers) and headers[i] else f"列{i + 1}"
        items.append((str(key), value))
    return items


def compact_row_numbers(row_numbers):
    nums = [n for n in row_numbers or [] if n is not None]
    if not nums:
        return ""
    if len(nums) == 1:
        return str(nums[0])
    nums_sorted = sorted(nums)
    if nums_sorted == list(range(nums_sorted[0], nums_sorted[-1] + 1)):
        return f"{nums_sorted[0]}-{nums_sorted[-1]}"
    return ",".join(str(n) for n in nums)


def is_non_required_header(header):
    """Return True when a column header explicitly marks the column as 非必填/PPT非必填."""
    if header is None:
        return False
    normalized = str(header).replace(" ", "").replace("\n", "").lower()
    markers = ["非必填", "(非必填)", "（非必填）", "ppt非必填", "非必填列"]
    return any(marker in normalized for marker in markers)


def select_summary_columns(headers, rows=None):
    """Select summary columns with the same global rule used for tables.

    Do not apply a priority list or max-column cap here. A matched summary block
    inside a detail PPTX must show the same summary fields as the standalone
    summary PPTX, except columns explicitly marked as 非必填/PPT非必填.
    """
    return [i for i, h in enumerate(headers or []) if not is_non_required_header(h)]


def draw_summary_rows_table(slide, summary_context, summary_area, theme, shape):
    x = summary_area.get("x_cm", 0.8)
    y = summary_area.get("y_cm", 1.15)
    w = summary_area.get("w_cm", 20)
    h = summary_area.get("h_cm", 1.0)
    title_font = theme.get("title_font", "Microsoft YaHei")
    body_font = theme.get("body_font", title_font)
    accent = ensure_rgb(theme.get("accent"), RGBColor(0x1B, 0x36, 0x5D))
    header_bg = ensure_rgb(theme.get("header_bg"), RGBColor(0x1B, 0x36, 0x5D))
    header_fg = ensure_rgb(theme.get("header_fg"), RGBColor(0xFF, 0xFF, 0xFF))
    text_color = RGBColor(0x33, 0x33, 0x33)

    rows = summary_context.get("rows") or []
    headers = summary_context.get("headers") or []
    row_numbers = summary_context.get("summary_row_numbers") or []
    row_label = compact_row_numbers(row_numbers)
    note = summary_context.get("match_note") or ""
    prefix = f"{summary_context.get('summary_sheet_name', '汇总')} 对应行: {row_label}"
    if note:
        prefix = f"{prefix} ({note})"

    col_indices = select_summary_columns(headers, rows)
    if not col_indices:
        return shape

    table_w = max(1.0, w - 0.36)
    available_table_h = max(0.45, h - 0.48)
    row_count = len(rows) + 1
    col_count = len(col_indices) + 1
    header_h = 0.34
    data_row_h = 0.28
    desired_table_h = header_h + max(1, len(rows)) * data_row_h
    table_h = min(max(0.45, desired_table_h), available_table_h)
    block_y = y
    summary_widths = [0.9]
    for ci in col_indices:
        header_text = headers[ci] if ci < len(headers) else f"列{ci + 1}"
        sample_values = [row[ci] if ci < len(row) else "" for row in rows[:3]]
        max_len = max([len(str(header_text))] + [len(str(v)) for v in sample_values])
        summary_widths.append(min(4.2, max(1.15, max_len * 0.18 + 0.55)))
    summary_widths, actual_table_w = fit_col_widths(summary_widths, table_w, col_count)
    table_x = x + 0.18
    table_y = block_y + 0.40
    add_text(slide, prefix, x + 0.18, block_y + 0.06, w - 0.36, 0.3,
             title_font, 7, accent, True, PP_ALIGN.LEFT)
    table_shape = slide.shapes.add_table(row_count, col_count, cm(table_x), cm(table_y), cm(actual_table_w), cm(table_h))
    table = table_shape.table

    if table_h < desired_table_h:
        scale = table_h / max(0.01, desired_table_h)
        actual_header_h = max(0.18, header_h * scale)
        actual_data_h = max(0.18, data_row_h * scale)
    else:
        actual_header_h = header_h
        actual_data_h = data_row_h
    for i in range(row_count):
        table.rows[i].height = cm(actual_header_h if i == 0 else actual_data_h)
        tr = table.rows[i]._tr
        trPr = tr.find(f'{{{NS_A}}}trPr')
        if trPr is None:
            trPr = etree.SubElement(tr, f'{{{NS_A}}}trPr')
        trPr.set('hRule', 'exact')

    for c, width in enumerate(summary_widths):
        table.columns[c].width = cm(width)

    font_size = 5
    border_rgb = ensure_rgb(theme.get("border_color"), RGBColor(0x00, 0x00, 0x00))
    border_hex = str(border_rgb) if isinstance(border_rgb, RGBColor) else "000000"
    set_cell_bg(table.cell(0, 0), header_bg)
    set_cell_text(table.cell(0, 0), "汇总行", body_font, font_size, True, header_fg, PP_ALIGN.CENTER, max_chars=16)
    set_cell_border(table.cell(0, 0), border_hex)
    for c, ci in enumerate(col_indices, start=1):
        set_cell_bg(table.cell(0, c), header_bg)
        set_cell_text(table.cell(0, c), headers[ci] if ci < len(headers) else f"列{ci + 1}", body_font, font_size, True, header_fg, PP_ALIGN.CENTER, max_chars=22)
        set_cell_border(table.cell(0, c), border_hex)

    for r, row in enumerate(rows, start=1):
        set_cell_bg(table.cell(r, 0), RGBColor(0xFF, 0xFF, 0xFF))
        label = row_numbers[r - 1] if r - 1 < len(row_numbers) else r
        set_cell_text(table.cell(r, 0), label, body_font, font_size, False, text_color, PP_ALIGN.CENTER, max_chars=10)
        set_cell_border(table.cell(r, 0), border_hex)
        for c, ci in enumerate(col_indices, start=1):
            set_cell_bg(table.cell(r, c), RGBColor(0xFF, 0xFF, 0xFF))
            val = row[ci] if ci < len(row) else ""
            set_cell_text(table.cell(r, c), val, body_font, font_size, False, text_color, PP_ALIGN.LEFT, max_chars=36)
            set_cell_border(table.cell(r, c), border_hex)
    return shape


def draw_summary_context(slide, summary_context, summary_area, theme):
    if not summary_context or not summary_area:
        return None

    if summary_context.get("rows"):
        return draw_summary_rows_table(slide, summary_context, summary_area, theme, None)

    # Single-row summary is rendered as a one-row table rather than inline text.
    single_row_context = dict(summary_context)
    single_row_context["rows"] = [summary_context.get("row", [])]
    single_row_context["summary_row_numbers"] = [summary_context.get("summary_row_number")]
    return draw_summary_rows_table(slide, single_row_context, summary_area, theme, None)


def fit_col_widths(col_widths, available_w, num_cols):
    """Fit column widths to the available page width for dense, full-width tables."""
    total_w = sum(col_widths) if col_widths else 0
    if total_w > 0:
        widths = list(col_widths)
    else:
        widths = [1.2] * max(num_cols, 1)
        total_w = sum(widths)

    if available_w <= 0:
        return widths, total_w

    if total_w > 0:
        scale = available_w / total_w
        widths = [w * scale for w in widths]
    else:
        widths = [available_w / max(num_cols, 1)] * max(num_cols, 1)
    return widths, available_w


def draw_table(slide, headers, rows, col_widths, table_area, font_size, theme):
    num_cols = len(headers)
    num_rows = len(rows) + 1
    if num_rows == 0 or num_cols == 0:
        return None

    tf_name = theme.get("title_font", "Microsoft YaHei")
    hbg_rgb = ensure_rgb(theme.get("header_bg"), RGBColor(0x1B, 0x36, 0x5D))
    hfg_rgb = ensure_rgb(theme.get("header_fg"), RGBColor(0xFF, 0xFF, 0xFF))
    even_rgb = ensure_rgb(theme.get("row_even"), RGBColor(0xF5, 0xF5, 0xF0))
    odd_rgb = ensure_rgb(theme.get("row_odd"), RGBColor(0xFF, 0xFF, 0xFF))
    hf_name = theme.get("header_font", tf_name)
    bf_name = theme.get("body_font", tf_name)

    table_left = table_area.get("x_cm", 0.8)
    table_top = table_area.get("y_cm", 1.5)
    table_available_w = table_area.get("w_cm", 20)
    table_available_h = table_area.get("h_cm", 10)

    col_widths_norm, actual_table_w = fit_col_widths(col_widths, table_available_w, num_cols)
    row_h = min(0.48, max(0.24, table_available_h / num_rows))
    table_h = min(row_h * num_rows, table_available_h)
    actual_left = table_left
    actual_top = table_top

    table_shape = slide.shapes.add_table(
        num_rows, num_cols,
        cm(actual_left), cm(actual_top),
        cm(actual_table_w), cm(table_h)
    )
    table = table_shape.table

    for i, w in enumerate(col_widths_norm):
        table.columns[i].width = cm(w)
    for i in range(num_rows):
        row = table.rows[i]
        row.height = cm(row_h)
        tr = row._tr
        trPr = tr.find(f'{{{NS_A}}}trPr')
        if trPr is None:
            trPr = etree.SubElement(tr, f'{{{NS_A}}}trPr')
        trPr.set('hRule', 'exact')

    border_rgb = ensure_rgb(theme.get("border_color"), RGBColor(0x00, 0x00, 0x00))
    border_hex = str(border_rgb) if isinstance(border_rgb, RGBColor) else "000000"

    for ci in range(num_cols):
        cell = table.cell(0, ci)
        set_cell_text(cell, headers[ci], hf_name, font_size, bold=True, color=hfg_rgb, align=PP_ALIGN.CENTER)
        set_cell_bg(cell, hbg_rgb)
        set_cell_border(cell, border_hex)

    for ri, row_data in enumerate(rows):
        bg = even_rgb if ri % 2 == 0 else odd_rgb
        for ci in range(num_cols):
            cell = table.cell(ri + 1, ci)
            val = row_data[ci] if ci < len(row_data) else None
            col_w_cm = col_widths_norm[ci] if ci < len(col_widths_norm) else 2
            cell_max_chars = max(6, int(col_w_cm / max(font_size * 0.003, 0.015)))
            body_font_size = max(font_size - 1, 5)
            set_cell_text(cell, val, bf_name, body_font_size,
                          color=RGBColor(0x33, 0x33, 0x33), max_chars=cell_max_chars)
            set_cell_bg(cell, bg)
            set_cell_border(cell, border_hex)

    return table_shape


def draw_image_grid(slide, images, image_area, image_grid, theme):
    if not images or not image_area:
        return

    cols = max(1, int((image_grid or {}).get("cols", 1)))
    rows = max(1, int((image_grid or {}).get("rows", math.ceil(len(images) / cols))))
    gap = float((image_grid or {}).get("gap_cm", 0.18))
    caption_h = float((image_grid or {}).get("caption_h_cm", 0.32))

    ia_x = image_area.get("x_cm", 0.8)
    ia_y = image_area.get("y_cm", 1.5)
    ia_w = image_area.get("w_cm", 8)
    ia_h = image_area.get("h_cm", 10)

    target_px = float((image_grid or {}).get("target_max_px", 300))
    target_cm = (target_px / 96.0) * 2.54 if target_px > 0 else None
    raw_cell_w = max(0.5, (ia_w - gap * (cols - 1)) / cols)
    raw_cell_h = max(0.5, (ia_h - gap * (rows - 1)) / rows)
    if target_cm:
        cell_w = min(raw_cell_w, target_cm + 0.45)
        cell_h = min(raw_cell_h, target_cm + caption_h + 0.45)
    else:
        cell_w = raw_cell_w
        cell_h = raw_cell_h
    grid_w = cell_w * cols + gap * (cols - 1)
    grid_h = cell_h * rows + gap * (rows - 1)
    grid_x = ia_x + max(0.0, (ia_w - grid_w) / 2)
    grid_y = ia_y + max(0.0, (ia_h - grid_h) / 2)
    pic_h_limit = max(0.3, cell_h - caption_h)
    fill_order = (image_grid or {}).get("fill_order", "column_major")

    bf_name = theme.get("body_font", theme.get("title_font", "Microsoft YaHei"))
    caption_color = RGBColor(0x55, 0x55, 0x55)

    for j, img_info in enumerate(images):
        img_file = img_info.get("file")
        if not img_file or not os.path.exists(img_file):
            continue
        if fill_order == "column_major":
            col_j = j // rows
            row_j = j % rows
        else:
            col_j = j % cols
            row_j = j // cols
        if row_j >= rows or col_j >= cols:
            break
        cell_x = grid_x + col_j * (cell_w + gap)
        cell_y = grid_y + row_j * (cell_h + gap)
        try:
            pil_img = Image.open(img_file)
            pw_px, ph_px = pil_img.size
            pw_cm = pw_px / 96 * 2.54
            ph_cm = ph_px / 96 * 2.54
            scale_limits = [(cell_w - 0.05) / max(pw_cm, 0.01), pic_h_limit / max(ph_cm, 0.01), 1.0]
            if target_cm:
                scale_limits.append(target_cm / max(pw_cm, ph_cm, 0.01))
            scale = min(scale_limits)
            final_w = max(0.1, pw_cm * scale)
            final_h = max(0.1, ph_cm * scale)
            x = cell_x + (cell_w - final_w) / 2
            y = cell_y + (pic_h_limit - final_h) / 2
            slide.shapes.add_picture(img_file, cm(x), cm(y), cm(final_w), cm(final_h))

            caption = img_info.get("caption") or f"图片-{j + 1}"
            add_text(slide, caption, cell_x, cell_y + pic_h_limit, cell_w, caption_h,
                     bf_name, 7, caption_color, False, PP_ALIGN.CENTER)
        except Exception as e:
            print(f"Warning: image {img_file}: {e}", file=sys.stderr)


def build_table_page(prs, page, theme, sw, sh):
    slide = add_blank_slide(prs)

    sheet_name = page.get("sheet_name", "")
    page_idx = page.get("page_index", 0)
    page_total = page.get("page_total", 1)
    font_size = page.get("font_size", 10)
    col_widths = page.get("col_widths", [])
    headers = page.get("headers", [])
    rows = page.get("rows", [])
    images = page.get("images", [])
    layout_mode = page.get("layout", "full-table")

    tc_rgb = ensure_rgb(theme.get("title_color"), RGBColor(0x1B, 0x36, 0x5D))
    tf_name = theme.get("title_font", "Microsoft YaHei")

    # 姣忛〉涓绘爣棰樺彧浣跨敤鍒嗛〉琛ㄦ牸鐨?sheet 鍚嶏紱椤电爜浣滀负寮卞寲杈呭姪淇℃伅鏄剧ず銆?
    title_text = page.get("title") or sheet_name
    add_text(slide, title_text, 0.8, 0.25, sw - 4.2, 0.8, tf_name, 18, tc_rgb, True)
    if page_total > 1:
        add_text(slide, f"{page_idx + 1}/{page_total}", sw - 3.2, 0.35, 2.4, 0.5,
                 tf_name, 9, RGBColor(0x88, 0x88, 0x88), False, PP_ALIGN.RIGHT)

    if page.get("summary_context") and page.get("summary_area"):
        draw_summary_context(slide, page.get("summary_context"), page.get("summary_area"), theme)

    default_table_area = {
        "x_cm": 0.8,
        "y_cm": 1.35,
        "w_cm": page.get("table_width_cm", sw - 1.6),
        "h_cm": sh - 1.35 - 0.6,
    }
    table_area = page.get("table_area") or default_table_area
    draw_table(slide, headers, rows, col_widths, table_area, font_size, theme)

    if layout_mode in ("left-table-right-images", "top-table-bottom-images") and images:
        draw_image_grid(slide, images, page.get("image_area"), page.get("image_grid"), theme)

    return slide


def build_summary_context_page(prs, page, theme, sw, sh):
    slide = add_blank_slide(prs)
    tc_rgb = ensure_rgb(theme.get("title_color"), RGBColor(0x1B, 0x36, 0x5D))
    tf_name = theme.get("title_font", "Microsoft YaHei")
    title_text = page.get("title") or f"{page.get('sheet_name', 'Sheet')} 汇总信息"
    add_text(slide, title_text, 0.8, 0.25, sw - 4.2, 0.8, tf_name, 18, tc_rgb, True)
    if page.get("page_total", 1) > 1:
        add_text(slide, f"{page.get('page_index', 0) + 1}/{page.get('page_total', 1)}", sw - 3.2, 0.35, 2.4, 0.5,
                 tf_name, 9, RGBColor(0x88, 0x88, 0x88), False, PP_ALIGN.RIGHT)
    draw_summary_context(slide, page.get("summary_context"), page.get("summary_area"), theme)
    return slide


def build_image_page(prs, page, theme, sw, sh):
    slide = add_blank_slide(prs)
    tc_rgb = ensure_rgb(theme.get("title_color"), RGBColor(0x1B, 0x36, 0x5D))
    tf_name = theme.get("title_font", "Microsoft YaHei")
    add_text(slide, page.get("sheet_name", "图片"), 0.8, 0.3, sw - 1.6, 0.9, tf_name, 18, tc_rgb, True)

    images = page.get("images", [])
    if page.get("image_area"):
        draw_image_grid(slide, images, page.get("image_area"), page.get("image_grid"), theme)
        return slide

    for img_info in images:
        img_file = img_info.get("file")
        if not img_file or not os.path.exists(img_file):
            continue
        try:
            x = img_info.get("x", 0.8)
            y = img_info.get("y", 1.5)
            w = img_info.get("width_cm", 5)
            h = img_info.get("height_cm", 4)
            slide.shapes.add_picture(img_file, cm(x), cm(y), cm(w), cm(h))
            caption = img_info.get("caption")
            if caption:
                add_text(slide, caption, x, y + h, w, 0.35, theme.get("body_font", tf_name), 7,
                         RGBColor(0x55, 0x55, 0x55), False, PP_ALIGN.CENTER)
        except Exception as e:
            print(f"Warning: image {img_file}: {e}", file=sys.stderr)

    return slide


def build_pptx(layout_plan, output_path, theme_path=None, template_path=None):
    global_theme = {}
    if theme_path and os.path.exists(theme_path):
        with open(theme_path, encoding="utf-8") as f:
            global_theme = json.load(f)
    global_theme_obj = make_theme(global_theme)

    template_path = template_path or layout_plan.get("template_file_path")
    template_used = bool(template_path and os.path.exists(template_path))

    default_orient = layout_plan.get("orientation", "portrait")
    if default_orient == "landscape":
        fallback_w, fallback_h = 33.867, 19.05
    else:
        fallback_w, fallback_h = SLIDE_W_CM, SLIDE_H_CM
    default_w, default_h = get_plan_canvas_dims(layout_plan, fallback_w, fallback_h)

    if template_used:
        prs = Presentation(template_path)
        fit_template_to_layout_canvas(prs, default_w, default_h)
        clear_template_slides(prs)
    else:
        prs = Presentation()
        prs.slide_width = cm(default_w)
        prs.slide_height = cm(default_h)

    for page in layout_plan.get("pages", []):
        page_type = page.get("type")
        ptheme = get_page_theme(page, global_theme_obj)
        psw, psh = get_page_dims(page, default_w, default_h)

        if abs(psw - default_w) > 0.1 or abs(psh - default_h) > 0.1:
            psw, psh = default_w, default_h

        if page_type == "cover":
            build_cover(prs, page, ptheme, psw, psh)
        elif page_type == "toc":
            build_toc(prs, page, ptheme, psw, psh)
        elif page_type == "summary":
            build_summary(prs, page, ptheme, psw, psh)
        elif page_type == "summary_context":
            build_summary_context_page(prs, page, ptheme, psw, psh)
        elif page_type == "table":
            build_table_page(prs, page, ptheme, psw, psh)
        elif page_type == "images":
            build_image_page(prs, page, ptheme, psw, psh)

    prs.save(output_path)
    fsize = os.path.getsize(output_path)
    print(f"鉁?PPTX saved: {output_path} ({fsize/1024:.0f}KB, {len(layout_plan.get('pages',[]))} pages)", file=sys.stderr)
    return output_path


def main():
    parser = argparse.ArgumentParser(description="Build PPTX from layout plan")
    parser.add_argument("input", help="Layout plan JSON")
    parser.add_argument("--output", "-o", required=True, help="Output .pptx path")
    parser.add_argument("--theme", help="Default theme JSON file")
    parser.add_argument("--template", help="PPTX template file to use as the actual presentation base")
    args = parser.parse_args()

    with open(args.input, encoding="utf-8") as f:
        plan = json.load(f)

    build_pptx(plan, args.output, args.theme, args.template)


if __name__ == "__main__":
    main()

