# -*- coding: utf-8 -*-
"""
差旅费用PPT生成脚本
用法: python 生成差旅PPT.py --excel 差旅表.xlsx --images 城际交通 --output 差旅结算.pptx
"""

import os
import json
import base64
import argparse
from openai import OpenAI
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from lxml import etree


# ============ VLM配置 ============
VLM_BASE_URL = "http://116.62.224.41:17600/v1"
VLM_API_KEY = "sk-K89PAzaakOKYY0wEZ9omLtW9jWyOhDwdyQftORQn7Ds2UFpm"
MODEL_NAME = "qwen3.6-27b"


def encode_image(image_path):
    """将图片编码为base64"""
    with open(image_path, 'rb') as f:
        return base64.b64encode(f.read()).decode('utf-8')


def recognize_invoice_vlm(image_path):
    """使用VLM识别发票图片，返回JSON格式结果"""
    
    client = OpenAI(
        base_url=VLM_BASE_URL,
        api_key=VLM_API_KEY
    )
    
    # 编码图片
    base64_image = encode_image(image_path)
    
    # 构建提示词
    prompt = """你是一个专业的发票识别助手。请从这张发票/行程单图片中提取以下信息，并以标准JSON格式输出：

提取字段：
- name: 乘客姓名（字符串）
- date: 出行日期（格式：YYYY-MM-DD）
- departure: 出发地（字符串，如机场或车站名称）
- destination: 目的地（字符串，如机场或车站名称）
- travel_type: 出行方式（字符串，如"飞机"、"高铁"、"火车"）
- amount: 金额（数字，保留两位小数）

注意事项：
1. 如果某个字段无法识别，设置为null
2. 金额只保留数字，不要包含货币符号
3. 日期必须是标准格式：YYYY-MM-DD
4. 只输出JSON，不要输出其他内容

请直接输出JSON："""
    
    # 调用VLM
    response = client.chat.completions.create(
        model=MODEL_NAME,
        messages=[
            {
                "role": "user",
                "content": [
                    {
                        "type": "text",
                        "text": prompt
                    },
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": f"data:image/png;base64,{base64_image}"
                        }
                    }
                ]
            }
        ],
        temperature=0.1
    )
    
    # 解析结果
    result_text = response.choices[0].message.content.strip()
    
    # 尝试解析JSON
    try:
        result = json.loads(result_text)
    except json.JSONDecodeError:
        # 去掉markdown代码块
        if "```json" in result_text:
            start = result_text.find("```json") + 7
            end = result_text.find("```", start)
            result_text = result_text[start:end].strip()
        elif "```" in result_text:
            start = result_text.find("```") + 3
            end = result_text.find("```", start)
            result_text = result_text[start:end].strip()
        
        try:
            result = json.loads(result_text)
        except json.JSONDecodeError:
            return None
    
    return result


def convert_vlm_to_info(vlm_result):
    """将VLM结果转换为内部格式"""
    if not vlm_result:
        return None
    
    info = {}
    
    # 姓名
    if vlm_result.get('name'):
        info['姓名'] = vlm_result['name']
    
    # 日期
    if vlm_result.get('date'):
        info['时间'] = vlm_result['date'][5:7] + '-' + vlm_result['date'][8:10]  # MM-DD格式
        info['日期'] = vlm_result['date']  # 完整日期
    
    # 行程
    if vlm_result.get('departure') and vlm_result.get('destination'):
        info['出发地'] = vlm_result['departure']
        info['目的地'] = vlm_result['destination']
        info['行程'] = f"{vlm_result['departure']}-{vlm_result['destination']}"
    
    # 金额
    if vlm_result.get('amount'):
        info['金额'] = float(vlm_result['amount'])
    
    # 出行方式
    if vlm_result.get('travel_type'):
        travel_type = vlm_result['travel_type']
        if '高铁' in travel_type or '火车' in travel_type:
            info['出行方式'] = '高铁'
        else:
            info['出行方式'] = '飞机'
    
    return info


def read_excel_summary(excel_path):
    """读取Excel中的差旅汇总信息"""
    wb = load_workbook(excel_path)
    ws = wb.active
    
    data = []
    for row in ws.iter_rows(min_row=3, values_only=True):
        if row[0] and '差旅' in str(row[0]):
            data.append({
                '板块': row[0],
                '描述方向': row[1],
                '媒体名称': row[2],
                '发布平台': row[3] if row[3] else '',
                '内容形式': row[4] if row[4] else '',
                '执行单价': row[8] if len(row) > 8 and row[8] else 0,
                '数量': row[9] if len(row) > 9 and row[9] else 1,
                '单位': row[10] if len(row) > 10 and row[10] else '项',
                '执行总价': row[11] if len(row) > 11 and row[11] else 0,
                '对应RFQ': row[12] if len(row) > 12 and row[12] else '',
                '备注': row[13] if len(row) > 13 and row[13] else ''
            })
            print(f"   读取: {row[1]}, 平台={row[3]}, 总价={row[11]}")
    
    return data


def set_cell_border(cell, border_color=RGBColor(0, 0, 0), border_width=Pt(1)):
    """设置单元格边框"""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    color_val = '%02X%02X%02X' % (border_color[0], border_color[1], border_color[2])
    for border_name in ['a:lnL', 'a:lnR', 'a:lnT', 'a:lnB']:
        ln = etree.SubElement(tcPr, qn(border_name))
        ln.set('w', str(int(border_width)))
        ln.set('cap', 'flat')
        ln.set('cmpd', 'sng')
        solidFill = etree.SubElement(ln, qn('a:solidFill'))
        srgbClr = etree.SubElement(solidFill, qn('a:srgbClr'))
        srgbClr.set('val', color_val)
        etree.SubElement(ln, qn('a:round'))


def create_pptx(summary_data, transport_data, output_path):
    """生成差旅结算PPT"""
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    title_color = RGBColor(0, 51, 102)

    total_amount = sum(item.get('金额', 0) or 0 for item in transport_data) if transport_data else 0
    summary = [s for s in summary_data if
               '交通' in str(s.get('发布平台', '')) or
               '机票' in str(s.get('发布平台', '')) or
               '高铁' in str(s.get('发布平台', '')) or
               ('/' in str(s.get('发布平台', '')) and '差旅' in str(s.get('板块', '')))]

    row_height = Inches(0.4)
    header_height = Inches(0.4)
    title_height = Inches(1.0)
    summary_table_height = len(summary) * row_height + header_height if summary else 0
    padding = Inches(0.3)

    available_height = prs.slide_height - title_height - summary_table_height - padding
    rows_per_page = max(1, int(available_height / row_height))

    slides = []
    current_slide = None
    current_detail_start = 0

    while current_detail_start < len(transport_data):
        current_slide = prs.slides.add_slide(prs.slide_layouts[6])
        slides.append(current_slide)

        title_box = current_slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12), Inches(0.6))
        tf = title_box.text_frame
        p = tf.paragraphs[0]
        p.text = f"人员差旅（城际交通）  ¥{total_amount:,.2f}元"
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = title_color

        current_top = title_height

        if summary and current_detail_start == 0:
            create_summary_table(current_slide, summary, Inches(0.5), current_top)
            current_top += summary_table_height + Inches(0.2)

        page_rows = transport_data[current_detail_start:current_detail_start + rows_per_page]
        detail_rows = len(page_rows) + 1
        detail_height = detail_rows * row_height

        create_detail_table_with_header(current_slide, page_rows, Inches(0.5), current_top, row_height)

        current_detail_start += rows_per_page

    for i in range(len(slides) - 1, -1, -1):
        slide = slides[i]
        has_content = False
        for shape in slide.shapes:
            if shape.has_table:
                has_content = True
                break
        if not has_content:
            rId = prs.slides._sldIdLst[i].rId
            prs.part.drop_rel(rId)
            del prs.slides._sldIdLst[i]

    if os.path.exists(output_path):
        try:
            os.remove(output_path)
        except:
            pass

    prs.save(output_path)

    if os.path.exists(output_path) and os.path.getsize(output_path) > 1000:
        print(f"PPT已生成: {output_path}")
    else:
        print(f"警告: PPT文件可能未正确生成")


def create_summary_table(slide, data, left, top):
    """创建汇总表格"""
    if not data:
        return

    all_keys = ['板块', '描述方向', '媒体名称', '发布平台', '内容形式', '执行单价', '数量', '单位', '执行总价', '对应RFQ', '备注']
    headers = [h for h in all_keys if any(item.get(h) for item in data)]
    if not headers:
        headers = ['内容方向', '媒体', '发布平台', '执行总价', '对应RFQ']

    rows = len(data) + 1
    cols = len(headers)
    table_width = Inches(12)
    slide_width = Inches(13.333)
    left = (slide_width - table_width) / 2 + Inches(1.1)
    top = top - Inches(0.3)
    row_height = Inches(0.4)
    table_height = rows * row_height

    col_width = table_width // cols
    font_size = Pt(10)
    if cols > 6:
        font_size = Pt(8)
    if cols > 8:
        font_size = Pt(7)

    table = slide.shapes.add_table(rows, cols, left, top, table_width, table_height).table

    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        set_cell_style(cell, True, font_size)
        set_cell_border(cell)

    for row_idx, item in enumerate(data, 1):
        for col_idx, key in enumerate(headers):
            val = item.get(key, '')
            if key == '执行单价':
                cell_text = f"¥{float(val or 0):,.2f}"
            elif key == '执行总价':
                cell_text = f"¥{float(val or 0):,.2f}"
            else:
                cell_text = str(val) if val else ''
            table.cell(row_idx, col_idx).text = cell_text

        for i in range(cols):
            set_cell_style(table.cell(row_idx, i), False, font_size)
            set_cell_border(table.cell(row_idx, i))


def create_detail_table_with_header(slide, data, left, top, row_height):
    """创建明细表格（带表头）"""
    if not data:
        return
    rows = len(data) + 1
    cols = 5
    table_width = Inches(12)
    slide_width = Inches(13.333)
    left = (slide_width - table_width) / 2 + Inches(1.1)
    top = top - Inches(0.3)
    table_height = rows * row_height

    table = slide.shapes.add_table(rows, cols, left, top, table_width, table_height).table

    headers = ['姓名', '时间', '行程', '出行方式', '金额']
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        set_cell_style(cell, True)
        set_cell_border(cell)

    for row_idx, item in enumerate(data, 1):
        table.cell(row_idx, 0).text = str(item.get('姓名', ''))
        table.cell(row_idx, 1).text = str(item.get('时间', ''))
        table.cell(row_idx, 2).text = str(item.get('行程', ''))
        table.cell(row_idx, 3).text = str(item.get('出行方式', ''))
        table.cell(row_idx, 4).text = f"¥{item.get('金额', 0) or 0:,.2f}"

        for i in range(cols):
            set_cell_style(table.cell(row_idx, i), False)
            set_cell_border(table.cell(row_idx, i))


def set_cell_style(cell, is_header, font_size=Pt(12)):
    """设置单元格样式"""
    cell.text_frame.paragraphs[0].font.size = font_size
    cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER


def main():
    parser = argparse.ArgumentParser(description='差旅费用PPT生成工具')
    parser.add_argument('--excel', '-e', default=r'差旅表.xlsx', help='Excel文件路径')
    parser.add_argument('--images', '-i', default=r'城际交通', help='发票图片文件夹路径')
    parser.add_argument('--output', '-o', default=r'差旅结算.pptx', help='输出PPT文件路径')
    args = parser.parse_args()
    
    excel_path = args.excel
    image_folder = args.images
    output_path = args.output
    
    print(f"Excel路径: {excel_path}")
    print(f"图片文件夹: {image_folder}")
    print(f"输出PPT: {output_path}")
    print()
    
    print("1. 读取Excel汇总数据...")
    summary_data = read_excel_summary(excel_path)
    print(f"   读取到 {len(summary_data)} 条汇总记录")
    
    print("2. 使用VLM识别发票图片...")
    transport_data = []
    
    images = sorted([f for f in os.listdir(image_folder) if f.endswith('.png')])
    
    for img_name in images:
        img_path = os.path.join(image_folder, img_name)
        print(f"   正在识别: {img_name}")
        
        vlm_result = recognize_invoice_vlm(img_path)
        
        if vlm_result:
            info = convert_vlm_to_info(vlm_result)
            if info and info.get('姓名'):
                transport_data.append(info)
                print(f"      -> {info.get('姓名')}, {info.get('行程')}, ¥{info.get('金额'):.2f}")
            else:
                print(f"      -> 解析失败: {vlm_result}")
        else:
            print(f"      -> VLM识别失败")
    
    print(f"   共识别到 {len(transport_data)} 条交通记录")
    
    print("3. 生成PPT...")
    create_pptx(summary_data, transport_data, output_path)
    
    print("\n完成！")


if __name__ == '__main__':
    main()
