#!/usr/bin/env python3
"""Audit and conservatively optimize PPTX layout overlaps.

This script is intentionally deterministic. It uses python-pptx geometry only,
so it never deletes content or rewrites text. It detects high-risk rectangular
shape overlaps and moves low-priority movable shapes to nearby safe positions.
"""

from __future__ import annotations

import argparse
import json
import math
import shutil
import sys
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Inches
except Exception as exc:  # pragma: no cover - shown to CLI users
    print(
        "Missing dependency: python-pptx. Install with: sudo pip3 install python-pptx",
        file=sys.stderr,
    )
    raise

EMU_PER_INCH = 914400

MODE_CONFIG = {
    "audit": {
        "threshold": 0.08,
        "max_iterations": 0,
        "grid_cols": 4,
        "grid_rows": 3,
        "move_titles": False,
        "move_placeholders": False,
    },
    "safe": {
        "threshold": 0.14,
        "max_iterations": 60,
        "grid_cols": 4,
        "grid_rows": 3,
        "move_titles": False,
        "move_placeholders": False,
    },
    "balanced": {
        "threshold": 0.08,
        "max_iterations": 120,
        "grid_cols": 5,
        "grid_rows": 4,
        "move_titles": False,
        "move_placeholders": True,
    },
    "aggressive": {
        "threshold": 0.035,
        "max_iterations": 180,
        "grid_cols": 6,
        "grid_rows": 5,
        "move_titles": False,
        "move_placeholders": True,
    },
}


@dataclass
class Rect:
    x: int
    y: int
    w: int
    h: int

    @property
    def right(self) -> int:
        return self.x + self.w

    @property
    def bottom(self) -> int:
        return self.y + self.h

    @property
    def area(self) -> int:
        return max(0, self.w) * max(0, self.h)

    def intersection_area(self, other: "Rect") -> int:
        ix = max(0, min(self.right, other.right) - max(self.x, other.x))
        iy = max(0, min(self.bottom, other.bottom) - max(self.y, other.y))
        return ix * iy

    def moved_to(self, x: int, y: int) -> "Rect":
        return Rect(int(x), int(y), self.w, self.h)

    def as_cm_tuple(self) -> Tuple[float, float, float, float]:
        return (
            round(self.x / 360000, 2),
            round(self.y / 360000, 2),
            round(self.w / 360000, 2),
            round(self.h / 360000, 2),
        )


@dataclass
class ShapeInfo:
    uid: str
    slide_index: int
    z_index: int
    shape: Any
    name: str
    shape_type: str
    category: str
    rect: Rect
    text: str = ""
    locked: bool = False
    skip: bool = False
    skip_reason: str = ""
    priority: int = 50
    shape_id: Optional[int] = None
    placeholder_type: str = ""
    notes: List[str] = field(default_factory=list)

    @property
    def area(self) -> int:
        return self.rect.area

    @property
    def has_text(self) -> bool:
        return bool(self.text.strip())

    @property
    def label(self) -> str:
        text = self.text.strip().replace("\n", " ")
        if len(text) > 34:
            text = text[:34] + "…"
        base = self.name or self.category
        if text:
            return f"{base} | {text}"
        return base


@dataclass
class Collision:
    slide_index: int
    a: ShapeInfo
    b: ShapeInfo
    intersection: int
    smaller_ratio: float
    larger_ratio: float
    severity: str

    @property
    def top_shape(self) -> ShapeInfo:
        return self.a if self.a.z_index > self.b.z_index else self.b

    @property
    def bottom_shape(self) -> ShapeInfo:
        return self.b if self.a.z_index > self.b.z_index else self.a

    def as_dict(self) -> Dict[str, Any]:
        return {
            "slide": self.slide_index,
            "shape_a": self.a.label,
            "shape_b": self.b.label,
            "shape_a_uid": self.a.uid,
            "shape_b_uid": self.b.uid,
            "top_shape_uid": self.top_shape.uid,
            "intersection_cm2": round(self.intersection / (360000 * 360000), 2),
            "smaller_ratio": round(self.smaller_ratio, 4),
            "larger_ratio": round(self.larger_ratio, 4),
            "severity": self.severity,
        }


@dataclass
class ActionRecord:
    slide_index: int
    shape_uid: str
    shape_label: str
    action: str
    before_cm: Tuple[float, float, float, float]
    after_cm: Tuple[float, float, float, float]
    reason: str

    def as_dict(self) -> Dict[str, Any]:
        return {
            "slide": self.slide_index,
            "shape_uid": self.shape_uid,
            "shape": self.shape_label,
            "action": self.action,
            "before_cm": self.before_cm,
            "after_cm": self.after_cm,
            "reason": self.reason,
        }


@dataclass
class SlideReport:
    slide_index: int
    shape_count: int = 0
    checked_shape_count: int = 0
    skipped_shape_count: int = 0
    initial_collisions: List[Collision] = field(default_factory=list)
    final_collisions: List[Collision] = field(default_factory=list)
    actions: List[ActionRecord] = field(default_factory=list)
    unresolved_notes: List[str] = field(default_factory=list)

    def as_dict(self) -> Dict[str, Any]:
        return {
            "slide": self.slide_index,
            "shape_count": self.shape_count,
            "checked_shape_count": self.checked_shape_count,
            "skipped_shape_count": self.skipped_shape_count,
            "initial_collision_count": len(self.initial_collisions),
            "final_collision_count": len(self.final_collisions),
            "actions": [a.as_dict() for a in self.actions],
            "final_collisions": [c.as_dict() for c in self.final_collisions],
            "unresolved_notes": self.unresolved_notes,
        }


def emu(value: float) -> int:
    return int(round(value * EMU_PER_INCH))


def safe_int(value: Any, default: int = 0) -> int:
    try:
        return int(value)
    except Exception:
        return default


def shape_text(shape: Any) -> str:
    try:
        if getattr(shape, "has_text_frame", False):
            return shape.text or ""
    except Exception:
        return ""
    return ""


def shape_type_name(shape: Any) -> str:
    try:
        return str(shape.shape_type)
    except Exception:
        return "unknown"


def placeholder_type_name(shape: Any) -> str:
    try:
        if getattr(shape, "is_placeholder", False):
            return str(shape.placeholder_format.type)
    except Exception:
        return ""
    return ""


def detect_category(shape: Any, rect: Rect, slide_w: int, slide_h: int, z_index: int) -> Tuple[str, bool, str, int, List[str]]:
    notes: List[str] = []
    st = shape_type_name(shape).lower()
    name = (getattr(shape, "name", "") or "").lower()
    text = shape_text(shape)
    ph_type = placeholder_type_name(shape).lower()
    slide_area = slide_w * slide_h
    area_ratio = rect.area / slide_area if slide_area else 0

    category = "shape"
    priority = 40
    skip = False
    skip_reason = ""

    if rect.w <= emu(0.03) or rect.h <= emu(0.03):
        return "decoration", True, "极细线条或尺寸过小，通常为装饰/连接线", 10, notes

    if "line" in st or "connector" in st:
        return "connector", True, "连接线暂不参与几何遮挡修复", 20, notes

    try:
        if getattr(shape, "has_table", False):
            category = "table"
            priority = 65
        elif getattr(shape, "has_chart", False):
            category = "chart"
            priority = 65
        elif "picture" in st:
            category = "picture"
            priority = 45
        elif "group" in st:
            category = "group"
            priority = 70
        elif text.strip():
            category = "text"
            priority = 55
        else:
            category = "shape"
            priority = 35
    except Exception:
        pass

    if area_ratio >= 0.78 and z_index <= 1:
        return "background", True, "底层大面积对象，按背景处理", 100, notes

    if "background" in name or "背景" in name:
        return "background", True, "名称显示为背景对象", 100, notes

    top_zone = rect.y < slide_h * 0.22
    bottom_zone = rect.bottom > slide_h * 0.90
    small_height = rect.h < slide_h * 0.15

    title_hint = (
        "title" in name
        or "标题" in name
        or "title" in ph_type
        or "center title" in ph_type
    )
    if title_hint and top_zone:
        return "title", False, "", 95, notes

    if ("footer" in name or "页脚" in name or "slide number" in ph_type) and bottom_zone:
        return "footer", False, "", 85, notes

    if (top_zone or bottom_zone) and small_height and category == "text" and len(text.strip()) < 60:
        category = "header_footer"
        priority = 80
        notes.append("边缘短文本，疑似页眉或页脚")

    return category, skip, skip_reason, priority, notes


def collect_slide_shapes(slide: Any, slide_index: int, slide_w: int, slide_h: int, mode: str) -> List[ShapeInfo]:
    infos: List[ShapeInfo] = []
    cfg = MODE_CONFIG[mode]
    for z_index, shape in enumerate(slide.shapes):
        try:
            rect = Rect(
                safe_int(shape.left),
                safe_int(shape.top),
                safe_int(shape.width),
                safe_int(shape.height),
            )
        except Exception:
            continue
        name = getattr(shape, "name", "") or f"shape-{z_index}"
        category, skip, skip_reason, priority, notes = detect_category(shape, rect, slide_w, slide_h, z_index)
        ph_type = placeholder_type_name(shape)
        locked = category in {"background", "title", "footer", "header_footer"}
        if category == "group" and mode == "safe":
            locked = True
            notes.append("safe 模式下锁定组合形状")
        if getattr(shape, "is_placeholder", False) and not cfg["move_placeholders"]:
            if category not in {"background", "title", "footer", "header_footer"}:
                locked = True
                notes.append("当前模式默认不移动普通占位符")
        if category == "title" and cfg["move_titles"] is False:
            locked = True
        info = ShapeInfo(
            uid=f"s{slide_index}-z{z_index}",
            slide_index=slide_index,
            z_index=z_index,
            shape=shape,
            name=name,
            shape_type=shape_type_name(shape),
            category=category,
            rect=rect,
            text=shape_text(shape),
            locked=locked,
            skip=skip,
            skip_reason=skip_reason,
            priority=priority,
            shape_id=getattr(shape, "shape_id", None),
            placeholder_type=ph_type,
            notes=notes,
        )
        infos.append(info)
    return infos


def severity_from_ratio(ratio: float) -> str:
    if ratio >= 0.45:
        return "high"
    if ratio >= 0.18:
        return "medium"
    return "low"


def detect_collisions(infos: Sequence[ShapeInfo], slide_area: int, threshold: float) -> List[Collision]:
    collisions: List[Collision] = []
    active = [i for i in infos if not i.skip and i.area > 0]
    for idx, a in enumerate(active):
        if a.category == "background":
            continue
        for b in active[idx + 1 :]:
            if b.category == "background":
                continue
            if a.category in {"connector", "decoration"} and b.category in {"connector", "decoration"}:
                continue
            inter = a.rect.intersection_area(b.rect)
            if inter <= 0:
                continue
            smaller = max(1, min(a.area, b.area))
            larger = max(1, max(a.area, b.area))
            smaller_ratio = inter / smaller
            larger_ratio = inter / larger
            local_threshold = threshold
            if a.has_text or b.has_text:
                local_threshold = min(local_threshold, 0.04)
            if (a.has_text and b.has_text):
                local_threshold = min(local_threshold, 0.025)
            if smaller_ratio < local_threshold:
                continue
            if inter < slide_area * 0.00025 and smaller_ratio < 0.20:
                continue
            collisions.append(
                Collision(
                    slide_index=a.slide_index,
                    a=a,
                    b=b,
                    intersection=inter,
                    smaller_ratio=smaller_ratio,
                    larger_ratio=larger_ratio,
                    severity=severity_from_ratio(smaller_ratio),
                )
            )
    collisions.sort(key=lambda c: (c.smaller_ratio, c.intersection), reverse=True)
    return collisions


def choose_movable(collision: Collision) -> Optional[ShapeInfo]:
    a, b = collision.a, collision.b
    candidates = [s for s in (a, b) if not s.locked and not s.skip]
    if not candidates:
        return None
    top = collision.top_shape
    if top in candidates and top.priority <= max(c.priority for c in candidates):
        return top
    candidates.sort(key=lambda s: (s.priority, -s.z_index))
    return candidates[0]


def clamp(value: int, low: int, high: int) -> int:
    return max(low, min(high, value))


def inside_slide(rect: Rect, slide_w: int, slide_h: int, margin: int) -> bool:
    return rect.x >= margin and rect.y >= margin and rect.right <= slide_w - margin and rect.bottom <= slide_h - margin


def candidate_positions(target: ShapeInfo, blocker: ShapeInfo, slide_w: int, slide_h: int, margin: int, gap: int, grid_cols: int, grid_rows: int) -> List[Rect]:
    r = target.rect
    o = blocker.rect
    candidates: List[Rect] = []

    direct = [
        (r.x, o.bottom + gap),
        (r.x, o.y - r.h - gap),
        (o.right + gap, r.y),
        (o.x - r.w - gap, r.y),
        (o.right + gap, o.bottom + gap),
        (o.x - r.w - gap, o.y - r.h - gap),
    ]
    for x, y in direct:
        candidates.append(r.moved_to(x, y))

    usable_w = max(1, slide_w - 2 * margin - r.w)
    usable_h = max(1, slide_h - 2 * margin - r.h)
    for row in range(grid_rows):
        for col in range(grid_cols):
            x = margin + int(round((usable_w * col) / max(1, grid_cols - 1)))
            y = margin + int(round((usable_h * row) / max(1, grid_rows - 1)))
            candidates.append(r.moved_to(x, y))

    clamped = []
    for c in candidates:
        if r.w <= slide_w - 2 * margin:
            x = clamp(c.x, margin, slide_w - margin - r.w)
        else:
            x = margin
        if r.h <= slide_h - 2 * margin:
            y = clamp(c.y, margin, slide_h - margin - r.h)
        else:
            y = margin
        clamped.append(r.moved_to(x, y))

    unique: Dict[Tuple[int, int], Rect] = {}
    for c in clamped:
        unique[(c.x, c.y)] = c
    return list(unique.values())


def score_rect(candidate: Rect, target: ShapeInfo, infos: Sequence[ShapeInfo], slide_w: int, slide_h: int, margin: int, threshold: float) -> float:
    if not inside_slide(candidate, slide_w, slide_h, margin):
        return float("inf")
    penalty = 0.0
    for other in infos:
        if other.uid == target.uid or other.skip or other.category == "background":
            continue
        inter = candidate.intersection_area(other.rect)
        if inter <= 0:
            continue
        smaller = max(1, min(candidate.area, other.area))
        ratio = inter / smaller
        local_threshold = threshold
        if target.has_text or other.has_text:
            local_threshold = min(local_threshold, 0.04)
        if target.has_text and other.has_text:
            local_threshold = min(local_threshold, 0.025)
        if ratio >= local_threshold:
            weight = 8.0 if other.locked else 2.0
            if target.has_text or other.has_text:
                weight *= 1.5
            penalty += ratio * weight + (inter / max(1, candidate.area))
    dx = abs(candidate.x - target.rect.x) / max(1, slide_w)
    dy = abs(candidate.y - target.rect.y) / max(1, slide_h)
    movement = dx + dy
    return penalty * 1000.0 + movement


def find_best_position(target: ShapeInfo, collision: Collision, infos: Sequence[ShapeInfo], slide_w: int, slide_h: int, mode: str, margin: int, gap: int) -> Tuple[Optional[Rect], float, float]:
    cfg = MODE_CONFIG[mode]
    blocker = collision.b if collision.a.uid == target.uid else collision.a
    current_score = score_rect(target.rect, target, infos, slide_w, slide_h, margin, cfg["threshold"])
    best_rect: Optional[Rect] = None
    best_score = current_score
    for cand in candidate_positions(
        target,
        blocker,
        slide_w,
        slide_h,
        margin,
        gap,
        cfg["grid_cols"],
        cfg["grid_rows"],
    ):
        score = score_rect(cand, target, infos, slide_w, slide_h, margin, cfg["threshold"])
        if score < best_score:
            best_score = score
            best_rect = cand
    return best_rect, current_score, best_score


def apply_move(info: ShapeInfo, new_rect: Rect) -> None:
    info.shape.left = int(new_rect.x)
    info.shape.top = int(new_rect.y)
    info.rect = new_rect


def optimize_slide(slide: Any, slide_index: int, slide_w: int, slide_h: int, mode: str, margin: int, gap: int) -> SlideReport:
    cfg = MODE_CONFIG[mode]
    infos = collect_slide_shapes(slide, slide_index, slide_w, slide_h, mode)
    slide_area = slide_w * slide_h
    report = SlideReport(slide_index=slide_index)
    report.shape_count = len(infos)
    report.checked_shape_count = len([i for i in infos if not i.skip])
    report.skipped_shape_count = len([i for i in infos if i.skip])
    report.initial_collisions = detect_collisions(infos, slide_area, cfg["threshold"])

    if mode == "audit" or not report.initial_collisions:
        report.final_collisions = list(report.initial_collisions)
        return report

    no_progress = 0
    seen_attempts: set[Tuple[str, str]] = set()
    for _ in range(cfg["max_iterations"]):
        collisions = detect_collisions(infos, slide_area, cfg["threshold"])
        if not collisions:
            break
        collision = collisions[0]
        movable = choose_movable(collision)
        if movable is None:
            note = f"第 {slide_index} 页：{collision.a.label} 与 {collision.b.label} 均被锁定，无法自动移动。"
            if note not in report.unresolved_notes:
                report.unresolved_notes.append(note)
            break
        attempt_key = (movable.uid, collision.a.uid + ":" + collision.b.uid)
        best_rect, current_score, best_score = find_best_position(movable, collision, infos, slide_w, slide_h, mode, margin, gap)
        if best_rect is None or best_score >= current_score - 0.01 or attempt_key in seen_attempts:
            seen_attempts.add(attempt_key)
            no_progress += 1
            note = f"第 {slide_index} 页：{movable.label} 未找到更优安全位置，保留原位。"
            if note not in report.unresolved_notes:
                report.unresolved_notes.append(note)
            if no_progress >= 5:
                break
            # Try the next collision on the following iteration by marking this pair as lower priority is avoided
            # to keep the algorithm deterministic and non-destructive.
            break
        before = movable.rect.as_cm_tuple()
        try:
            apply_move(movable, best_rect)
            after = movable.rect.as_cm_tuple()
            report.actions.append(
                ActionRecord(
                    slide_index=slide_index,
                    shape_uid=movable.uid,
                    shape_label=movable.label,
                    action="move",
                    before_cm=before,
                    after_cm=after,
                    reason=f"减少与 {collision.b.label if collision.a.uid == movable.uid else collision.a.label} 的遮挡",
                )
            )
            seen_attempts.add(attempt_key)
            no_progress = 0
        except Exception as exc:
            note = f"第 {slide_index} 页：移动 {movable.label} 失败：{exc}"
            if note not in report.unresolved_notes:
                report.unresolved_notes.append(note)
            break

    report.final_collisions = detect_collisions(infos, slide_area, cfg["threshold"])
    return report


def write_markdown_report(path: Path, input_path: Path, output_path: Optional[Path], mode: str, slide_reports: Sequence[SlideReport]) -> None:
    total_initial = sum(len(r.initial_collisions) for r in slide_reports)
    total_final = sum(len(r.final_collisions) for r in slide_reports)
    total_actions = sum(len(r.actions) for r in slide_reports)
    risk = "低"
    if total_final >= 8:
        risk = "高"
    elif total_final >= 3:
        risk = "中"

    lines: List[str] = []
    lines.append("# PPTX 布局优化报告")
    lines.append("")
    lines.append(f"输入文件：`{input_path}`")
    if output_path:
        lines.append(f"输出文件：`{output_path}`")
    lines.append(f"运行模式：`{mode}`")
    lines.append("")
    lines.append("## 总览")
    lines.append("")
    lines.append("| 指标 | 数值 |")
    lines.append("| --- | ---: |")
    lines.append(f"| 幻灯片页数 | {len(slide_reports)} |")
    lines.append(f"| 初始遮挡风险 | {total_initial} |")
    lines.append(f"| 自动修复动作 | {total_actions} |")
    lines.append(f"| 剩余遮挡风险 | {total_final} |")
    lines.append(f"| 当前风险等级 | {risk} |")
    lines.append("")
    lines.append("## 分页摘要")
    lines.append("")
    lines.append("| 页码 | 检查对象 | 跳过对象 | 初始风险 | 修复动作 | 剩余风险 |")
    lines.append("| ---: | ---: | ---: | ---: | ---: | ---: |")
    for r in slide_reports:
        lines.append(
            f"| {r.slide_index} | {r.checked_shape_count} | {r.skipped_shape_count} | {len(r.initial_collisions)} | {len(r.actions)} | {len(r.final_collisions)} |"
        )
    lines.append("")

    if total_actions:
        lines.append("## 自动修复记录")
        lines.append("")
        lines.append("| 页码 | 元素 | 动作 | 原位置 cm `(x,y,w,h)` | 新位置 cm `(x,y,w,h)` | 原因 |")
        lines.append("| ---: | --- | --- | --- | --- | --- |")
        for r in slide_reports:
            for action in r.actions:
                lines.append(
                    f"| {action.slide_index} | {escape_md(action.shape_label)} | {action.action} | `{action.before_cm}` | `{action.after_cm}` | {escape_md(action.reason)} |"
                )
        lines.append("")

    if total_final:
        lines.append("## 剩余遮挡风险")
        lines.append("")
        lines.append("| 页码 | 元素 A | 元素 B | 严重程度 | 较小元素被覆盖比例 |")
        lines.append("| ---: | --- | --- | --- | ---: |")
        for r in slide_reports:
            for c in r.final_collisions:
                lines.append(
                    f"| {r.slide_index} | {escape_md(c.a.label)} | {escape_md(c.b.label)} | {c.severity} | {c.smaller_ratio:.1%} |"
                )
        lines.append("")
        lines.append("> 说明：剩余风险通常来自标题/页脚/背景等锁定元素、组合形状、空间不足，或需要人工重新设计的复杂版式。")
        lines.append("")

    notes = [note for r in slide_reports for note in r.unresolved_notes]
    if notes:
        lines.append("## 未自动处理说明")
        lines.append("")
        for note in notes:
            lines.append(f"- {escape_md(note)}")
        lines.append("")

    path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def escape_md(value: str) -> str:
    return value.replace("|", "\\|").replace("\n", " ")


def write_json_report(path: Path, input_path: Path, output_path: Optional[Path], mode: str, slide_reports: Sequence[SlideReport]) -> None:
    payload = {
        "input": str(input_path),
        "output": str(output_path) if output_path else None,
        "mode": mode,
        "summary": {
            "slides": len(slide_reports),
            "initial_collision_count": sum(len(r.initial_collisions) for r in slide_reports),
            "final_collision_count": sum(len(r.final_collisions) for r in slide_reports),
            "action_count": sum(len(r.actions) for r in slide_reports),
        },
        "slides": [r.as_dict() for r in slide_reports],
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def default_output_path(input_path: Path, mode: str) -> Path:
    stem = input_path.stem
    suffix = input_path.suffix or ".pptx"
    if mode == "audit":
        return input_path.with_name(f"{stem}_audit_copy{suffix}")
    return input_path.with_name(f"{stem}_layout_optimized{suffix}")


def parse_args(argv: Optional[Sequence[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Audit and optimize PPTX element overlaps without deleting content.")
    parser.add_argument("input", type=Path, help="Input .pptx file")
    parser.add_argument("-o", "--output", type=Path, help="Output optimized .pptx file")
    parser.add_argument("--report", type=Path, help="Markdown report path")
    parser.add_argument("--json-report", type=Path, help="JSON report path")
    parser.add_argument("--mode", choices=sorted(MODE_CONFIG.keys()), default="balanced", help="Optimization mode")
    parser.add_argument("--margin", type=float, default=0.22, help="Safe slide margin in inches")
    parser.add_argument("--gap", type=float, default=0.06, help="Minimum gap between separated shapes in inches")
    return parser.parse_args(argv)


def main(argv: Optional[Sequence[str]] = None) -> int:
    args = parse_args(argv)
    input_path: Path = args.input.expanduser().resolve()
    if not input_path.exists():
        print(f"Input file not found: {input_path}", file=sys.stderr)
        return 2
    if input_path.suffix.lower() not in {".pptx", ".pptm"}:
        print("Only .pptx/.pptm files are supported.", file=sys.stderr)
        return 2

    output_path = args.output.expanduser().resolve() if args.output else default_output_path(input_path, args.mode).resolve()
    report_path = args.report.expanduser().resolve() if args.report else output_path.with_suffix(".layout_report.md")
    json_report_path = args.json_report.expanduser().resolve() if args.json_report else output_path.with_suffix(".layout_report.json")

    prs = Presentation(str(input_path))
    slide_w, slide_h = int(prs.slide_width), int(prs.slide_height)
    margin = emu(args.margin)
    gap = emu(args.gap)

    slide_reports: List[SlideReport] = []
    for idx, slide in enumerate(prs.slides, start=1):
        slide_reports.append(optimize_slide(slide, idx, slide_w, slide_h, args.mode, margin, gap))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.parent.mkdir(parents=True, exist_ok=True)
    json_report_path.parent.mkdir(parents=True, exist_ok=True)

    if args.mode == "audit":
        if args.output:
            shutil.copyfile(input_path, output_path)
        else:
            output_path = None  # type: ignore[assignment]
    else:
        prs.save(str(output_path))

    write_markdown_report(report_path, input_path, output_path, args.mode, slide_reports)
    write_json_report(json_report_path, input_path, output_path, args.mode, slide_reports)

    total_initial = sum(len(r.initial_collisions) for r in slide_reports)
    total_final = sum(len(r.final_collisions) for r in slide_reports)
    total_actions = sum(len(r.actions) for r in slide_reports)
    print(f"Slides checked: {len(slide_reports)}")
    print(f"Initial overlap risks: {total_initial}")
    print(f"Actions applied: {total_actions}")
    print(f"Remaining overlap risks: {total_final}")
    if output_path:
        print(f"Output PPTX: {output_path}")
    print(f"Markdown report: {report_path}")
    print(f"JSON report: {json_report_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
